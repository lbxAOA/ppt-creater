from __future__ import annotations

import json
import subprocess
import sys
from pathlib import Path

from pptx import Presentation


def test_indexer_catalogs_slots_and_pptx_structure(tmp_path: Path) -> None:
    library = tmp_path / "library"
    library.mkdir()
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.shapes.add_textbox(0, 0, 1000000, 500000).name = "slot_title"
    deck.save(library / "template.pptx")
    output = tmp_path / "catalog.jsonl"
    script = Path(__file__).parents[1] / "src" / "ppt_creater" / "index_templates.py"
    completed = subprocess.run(
        [sys.executable, str(script), str(library), "--output", str(output)],
        capture_output=True,
        text=True,
        check=True,
    )
    assert '"files_scanned": 1' in completed.stdout
    record = json.loads(output.read_text(encoding="utf-8").strip())
    assert record["status"] == "ok"
    assert record["has_named_slots"] is True
    assert record["slots"][0]["shape_name"] == "slot_title"
