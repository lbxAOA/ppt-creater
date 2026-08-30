#!/usr/bin/env python
"""Enrich a PPTX catalog with text, theme tokens, and per-slide pattern signals."""
from __future__ import annotations

import argparse
import json
import re
import zipfile
from collections import Counter
from pathlib import Path
from typing import Any
from xml.etree import ElementTree as ET

from pptx import Presentation

NS = {
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
}


def theme_tokens(path: Path) -> dict[str, Any]:
    result: dict[str, Any] = {"colors": [], "fonts": []}
    try:
        with zipfile.ZipFile(path) as archive:
            theme_parts = [name for name in archive.namelist() if name.startswith("ppt/theme/") and name.endswith(".xml")]
            if not theme_parts:
                return result
            root = ET.fromstring(archive.read(theme_parts[0]))
            colors = []
            for node in root.findall(".//a:clrScheme/*", NS):
                value_node = next(iter(node), None)
                if value_node is not None:
                    value = value_node.attrib.get("lastClr") or value_node.attrib.get("val")
                    if value:
                        colors.append(value.upper())
            result["colors"] = list(dict.fromkeys(colors))[:20]
            major = root.find(".//a:themeElements/a:fontScheme/a:majorFont", NS)
            minor = root.find(".//a:themeElements/a:fontScheme/a:minorFont", NS)
            fonts = []
            for node in (major, minor):
                if node is None:
                    continue
                for typeface in node.findall(".//*[@typeface]", NS):
                    value = typeface.attrib.get("typeface")
                    if value and value not in fonts:
                        fonts.append(value)
            result["fonts"] = fonts[:20]
    except Exception as exc:
        result["theme_error"] = f"{type(exc).__name__}: {exc}"
    return result


def detect_pattern(slide: Any) -> str:
    text = " ".join(shape.text for shape in slide.shapes if getattr(shape, "has_text_frame", False)).lower()
    shape_count = len(slide.shapes)
    charts = sum(1 for shape in slide.shapes if getattr(shape, "has_chart", False))
    tables = sum(1 for shape in slide.shapes if getattr(shape, "has_table", False))
    if charts:
        return "chart"
    if tables:
        return "table"
    if any(word in text for word in ("目录", "contents", "agenda")):
        return "agenda"
    if any(word in text for word in ("谢谢", "thanks", "q&a", "thank you")):
        return "closing"
    if shape_count <= 4 and len(text) < 100:
        return "minimal"
    if shape_count >= 18:
        return "dense-composition"
    return "content"


def enrich(record: dict[str, Any]) -> dict[str, Any]:
    if record.get("status") != "ok":
        return record
    path = Path(record["absolute_path"])
    try:
        presentation = Presentation(path)
        slides = []
        full_text = []
        for index, slide in enumerate(presentation.slides, start=1):
            text_parts = [shape.text.strip() for shape in slide.shapes if getattr(shape, "has_text_frame", False) and shape.text.strip()]
            slide_text = "\n".join(text_parts)
            full_text.append(slide_text)
            slides.append({
                "slide_index": index,
                "pattern": detect_pattern(slide),
                "text_chars": len(slide_text),
                "text_excerpt": slide_text[:1200],
                "shape_count": len(slide.shapes),
            })
        record["theme"] = theme_tokens(path)
        record["slide_patterns"] = dict(Counter(slide["pattern"] for slide in slides))
        record["slides_enriched"] = slides
        record["text_excerpt"] = "\n---\n".join(full_text)[:10000]
        record["read_status"] = "complete"
    except Exception as exc:
        record["read_status"] = "error"
        record["read_error"] = f"{type(exc).__name__}: {exc}"
    return record


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("catalog", type=Path)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    records = [json.loads(line) for line in args.catalog.read_text(encoding="utf-8").splitlines() if line.strip()]
    args.output.parent.mkdir(parents=True, exist_ok=True)
    with args.output.open("w", encoding="utf-8") as handle:
        for i, record in enumerate(records, 1):
            handle.write(json.dumps(enrich(record), ensure_ascii=False) + "\n")
            if i % 50 == 0:
                print(json.dumps({"progress": i, "total": len(records)}, ensure_ascii=False), flush=True)
    print(json.dumps({"complete": True, "records": len(records)}, ensure_ascii=False))

if __name__ == "__main__":
    main()
