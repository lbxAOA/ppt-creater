from pathlib import Path
import json
import sys
sys.path.insert(0, r'C:/Users/rdft1/AppData/Local/hermes/skills/productivity/powerpoint/scripts')
from pptx import Presentation
for raw in sys.argv[1:]:
    p=Path(raw); prs=Presentation(p)
    rows=[]
    for slide in prs.slides:
        texts=[]
        for shape in slide.shapes:
            if getattr(shape,'has_text_frame',False) and shape.text.strip(): texts.append(shape.text.strip().replace('\n',' ')[:90])
        rows.append(texts)
    print(json.dumps({'file':str(p),'slides':len(prs.slides),'layouts':[x.name for x in prs.slide_layouts],'samples':rows[:5]},ensure_ascii=False))
