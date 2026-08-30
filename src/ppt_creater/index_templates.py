#!/usr/bin/env python
"""Build a non-destructive JSONL catalog for a local PPTX template library."""
from __future__ import annotations

import argparse
import hashlib
import json
import zipfile
from collections import Counter
from pathlib import Path
from typing import Any

from pptx import Presentation

NS = {"p": "http://schemas.openxmlformats.org/presentationml/2006/main"}


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def pptx_timing_summary(path: Path) -> dict[str, Any]:
    with zipfile.ZipFile(path) as archive:
        members = set(archive.namelist())
        timings = []
        transitions = []
        for name in sorted(member for member in members if member.startswith("ppt/slides/slide") and member.endswith(".xml")):
            data = archive.read(name)
            timings.append({"part": name, "has_timing": b"<p:timing" in data})
            transitions.append({"part": name, "has_transition": b"<p:transition" in data})
    return {
        "animated_slide_parts": [entry["part"] for entry in timings if entry["has_timing"]],
        "transition_slide_parts": [entry["part"] for entry in transitions if entry["has_transition"]],
    }


def shape_kind(shape: Any) -> str:
    if getattr(shape, "has_text_frame", False):
        return "text"
    if getattr(shape, "has_chart", False):
        return "chart"
    if getattr(shape, "has_table", False):
        return "table"
    if getattr(shape, "shape_type", None) is not None:
        return str(shape.shape_type)
    return "other"


def inspect_pptx(path: Path, library_root: Path) -> dict[str, Any]:
    record: dict[str, Any] = {
        "path": str(path),
        "relative_path": str(path.relative_to(library_root)),
        "bytes": path.stat().st_size,
        "sha256": sha256(path),
        "status": "ok",
    }
    try:
        presentation = Presentation(path)
        kinds: Counter[str] = Counter()
        slots: list[dict[str, Any]] = []
        slides: list[dict[str, Any]] = []
        for index, slide in enumerate(presentation.slides, start=1):
            shape_count = 0
            slide_slots: list[str] = []
            for shape in slide.shapes:
                shape_count += 1
                kinds[shape_kind(shape)] += 1
                name = str(getattr(shape, "name", ""))
                if name.startswith("slot_"):
                    slide_slots.append(name)
                    slots.append({"slide_index": index, "shape_name": name, "kind": shape_kind(shape)})
            slides.append({"slide_index": index, "shape_count": shape_count, "slot_names": slide_slots})
        timing = pptx_timing_summary(path)
        record.update({
            "slide_count": len(presentation.slides),
            "slide_size_inches": {
                "width": round(presentation.slide_width / 914400, 4),
                "height": round(presentation.slide_height / 914400, 4),
            },
            "layouts": [layout.name for layout in presentation.slide_layouts],
            "shape_kinds": dict(kinds),
            "slides": slides,
            "slots": slots,
            "animation": timing,
            "has_named_slots": bool(slots),
            "has_animation_xml": bool(timing["animated_slide_parts"]),
        })
    except Exception as exc:  # cataloging must continue past bad files
        record.update({"status": "error", "error": f"{type(exc).__name__}: {exc}"})
    return record


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("library_root", type=Path)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--limit", type=int, default=0)
    args = parser.parse_args()
    root = args.library_root.resolve()
    templates = sorted(path for path in root.rglob("*.pptx") if path.is_file() and not path.name.startswith("~$"))
    if args.limit:
        templates = templates[: args.limit]
    args.output.parent.mkdir(parents=True, exist_ok=True)
    with args.output.open("w", encoding="utf-8") as out:
        for path in templates:
            out.write(json.dumps(inspect_pptx(path, root), ensure_ascii=False) + "\n")
    print(json.dumps({"catalog": str(args.output), "files_scanned": len(templates)}, ensure_ascii=False))


if __name__ == "__main__":
    main()
