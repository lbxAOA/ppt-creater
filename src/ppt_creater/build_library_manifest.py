#!/usr/bin/env python
"""Create a non-destructive manifest and selection candidates for PPT/PPTX libraries."""
from __future__ import annotations

import argparse
import hashlib
import json
import re
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any

PRESENTATION_EXTENSIONS = {".pptx", ".ppt"}


def sha256(path: Path) -> str:
    hasher = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            hasher.update(chunk)
    return hasher.hexdigest()


def category(path: Path, root: Path) -> str:
    parts = path.relative_to(root).parts
    return parts[0] if parts else "unclassified"


def quality_score(item: dict[str, Any]) -> float:
    score = 0.0
    if item["format"] == ".pptx":
        score += 25
    size = item["bytes"]
    if 200_000 <= size <= 150_000_000:
        score += 10
    if item["format"] == ".pptx":
        if item.get("slide_count", 0) >= 8:
            score += 15
        elif item.get("slide_count", 0) >= 4:
            score += 8
        if item.get("has_animation_xml"):
            score += 15
        if item.get("has_named_slots"):
            score += 20
        score += min(item.get("layout_count", 0), 12) * 0.5
        score += min(item.get("shape_count", 0) / max(item.get("slide_count", 1), 1), 30) * 0.3
    return round(score, 2)


def inspect_pptx(path: Path) -> dict[str, Any]:
    from pptx import Presentation
    import zipfile

    data: dict[str, Any] = {}
    presentation = Presentation(path)
    slots: list[str] = []
    shape_count = 0
    chart_count = 0
    table_count = 0
    text_count = 0
    for slide in presentation.slides:
        for shape in slide.shapes:
            shape_count += 1
            name = str(getattr(shape, "name", ""))
            if name.startswith("slot_"):
                slots.append(name)
            if getattr(shape, "has_chart", False):
                chart_count += 1
            if getattr(shape, "has_table", False):
                table_count += 1
            if getattr(shape, "has_text_frame", False):
                text_count += 1
    with zipfile.ZipFile(path) as archive:
        slide_parts = [name for name in archive.namelist() if re.fullmatch(r"ppt/slides/slide\d+\.xml", name)]
        animated = [name for name in slide_parts if b"<p:timing" in archive.read(name)]
        transitions = [name for name in slide_parts if b"<p:transition" in archive.read(name)]
    data.update({
        "status": "ok",
        "slide_count": len(presentation.slides),
        "layout_count": len(presentation.slide_layouts),
        "slide_size_inches": {
            "width": round(presentation.slide_width / 914400, 4),
            "height": round(presentation.slide_height / 914400, 4),
        },
        "shape_count": shape_count,
        "text_count": text_count,
        "chart_count": chart_count,
        "table_count": table_count,
        "slot_names": slots,
        "has_named_slots": bool(slots),
        "has_animation_xml": bool(animated),
        "animated_slide_parts": animated,
        "has_transition_xml": bool(transitions),
        "transition_slide_parts": transitions,
    })
    return data


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("library_root", type=Path)
    parser.add_argument("--manifest", type=Path, required=True)
    parser.add_argument("--pptx-catalog", type=Path, required=True)
    parser.add_argument("--summary", type=Path, required=True)
    args = parser.parse_args()
    root = args.library_root.resolve()
    files = sorted(file for file in root.rglob("*") if file.is_file() and not file.name.startswith("~$"))
    all_records: list[dict[str, Any]] = []
    pptx_records: list[dict[str, Any]] = []
    errors: list[dict[str, str]] = []
    for index, path in enumerate(files, start=1):
        suffix = path.suffix.lower()
        record: dict[str, Any] = {
            "relative_path": str(path.relative_to(root)),
            "absolute_path": str(path),
            "category": category(path, root),
            "format": suffix or "[none]",
            "bytes": path.stat().st_size,
        }
        if suffix in PRESENTATION_EXTENSIONS:
            record["sha256"] = sha256(path)
        if suffix == ".pptx":
            try:
                record.update(inspect_pptx(path))
                record["quality_score"] = quality_score(record)
                pptx_records.append(record)
            except Exception as exc:
                record.update({"status": "error", "error": f"{type(exc).__name__}: {exc}", "quality_score": 0})
                errors.append({"path": record["relative_path"], "error": record["error"]})
        elif suffix == ".ppt":
            record.update({"status": "legacy_unconverted", "quality_score": quality_score(record)})
        all_records.append(record)
        if index % 100 == 0:
            print(json.dumps({"progress": index, "total": len(files)}, ensure_ascii=False), flush=True)

    exact_groups: dict[str, list[str]] = defaultdict(list)
    for record in all_records:
        if "sha256" in record:
            exact_groups[record["sha256"]].append(record["relative_path"])
    duplicates = [paths for paths in exact_groups.values() if len(paths) > 1]
    category_counts = Counter(record["category"] for record in all_records if record["format"] in PRESENTATION_EXTENSIONS)
    ranked = sorted((record for record in pptx_records if record["status"] == "ok"), key=lambda item: item["quality_score"], reverse=True)
    summary = {
        "library_root": str(root),
        "files_total": len(all_records),
        "presentations_total": sum(1 for record in all_records if record["format"] in PRESENTATION_EXTENSIONS),
        "pptx_total": sum(1 for record in all_records if record["format"] == ".pptx"),
        "ppt_total": sum(1 for record in all_records if record["format"] == ".ppt"),
        "pptx_indexed_ok": sum(1 for record in pptx_records if record["status"] == "ok"),
        "pptx_errors": errors,
        "animated_pptx": sum(1 for record in pptx_records if record.get("has_animation_xml")),
        "transition_pptx": sum(1 for record in pptx_records if record.get("has_transition_xml")),
        "slot_ready_pptx": sum(1 for record in pptx_records if record.get("has_named_slots")),
        "exact_duplicate_groups": duplicates,
        "category_counts": dict(category_counts),
        "top_candidates": [{key: record.get(key) for key in ("relative_path", "quality_score", "slide_count", "has_animation_xml", "has_named_slots", "chart_count", "table_count")} for record in ranked[:100]],
    }
    for output in (args.manifest, args.pptx_catalog, args.summary):
        output.parent.mkdir(parents=True, exist_ok=True)
    with args.manifest.open("w", encoding="utf-8") as handle:
        for record in all_records:
            handle.write(json.dumps(record, ensure_ascii=False) + "\n")
    with args.pptx_catalog.open("w", encoding="utf-8") as handle:
        for record in pptx_records:
            handle.write(json.dumps(record, ensure_ascii=False) + "\n")
    args.summary.write_text(json.dumps(summary, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"complete": True, "pptx_indexed": len(pptx_records), "errors": len(errors)}, ensure_ascii=False))


if __name__ == "__main__":
    main()
